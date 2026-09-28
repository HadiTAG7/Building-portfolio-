#!/usr/bin/env python3
"""Build tests/fixtures/excel-recalc.json from a *recalculated* copy of the workbook.

The source workbook ships without cached formula results, so it has to be
recalculated first (LibreOffice headless, see scripts/recalc_with_libreoffice.py).
The fixture captures the spreadsheet's own outputs so the TypeScript engine can
be checked against the Excel formulas.

Usage:
    python3 scripts/make_test_fixtures.py path/to/recalculated.xlsx
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tests" / "fixtures" / "excel-recalc.json"

METRIC_KEYS = [
    "days", "cumulative", "cagr", "avgDaily", "volatility", "sharpe",
    "sortino", "bestDay", "worstDay", "maxDrawdown", "downsideDeviation",
]


def sheet(wb, name):
    for ws in wb.worksheets:
        if ws.title.strip() == name.strip():
            return ws
    raise KeyError(name)


def iso(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)):  # Excel serial date
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    raise TypeError(v)


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def metric_blocks(ws, first_rows, n_portfolios):
    out = []
    for p in range(n_portfolios):
        per_window = []
        for r0 in first_rows:
            per_window.append({k: num(ws.cell(r0 + i, 2 + p).value) for i, k in enumerate(METRIC_KEYS)})
        out.append(per_window)
    return out


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ug, gr, rd = sheet(wb, "Ultra Growth"), sheet(wb, "Growth"), sheet(wb, "Risk Dashboard")
    summary, scen = sheet(wb, "Summary"), sheet(wb, "Scenario")
    tickers = [summary.cell(4 + i, 1).value for i in range(18)]

    windows = {
        "ultraGrowth": [
            {"start": None, "end": None},
            {"start": iso(ug["H27"].value), "end": iso(ug["J27"].value)},
            {"start": iso(ug["H41"].value), "end": iso(ug["J41"].value)},
        ],
        "growth": [
            {"start": None, "end": None},
            {"start": iso(gr["H26"].value), "end": iso(gr["J26"].value)},
            {"start": iso(gr["H40"].value), "end": iso(gr["J40"].value)},
        ],
    }
    portfolios = []
    for k, m in enumerate(metric_blocks(ug, [15, 29, 43], 7)):
        portfolios.append({"id": f"ug-{k + 1}", "group": "ultraGrowth", "metrics": m})
    for k, m in enumerate(metric_blocks(gr, [14, 28, 42], 6)):
        portfolios.append({"id": f"g-{k + 1}", "group": "growth", "metrics": m})

    corr = [[num(rd.cell(6 + i, 2 + j).value) for j in range(18)] for i in range(18)]

    drawdowns = []
    for i in range(5):
        drawdowns.append({"id": f"ug-{i + 1}", "maxDrawdown": num(rd.cell(26 + i, 2).value),
                          "troughDate": iso(rd.cell(26 + i, 3).value)})
        drawdowns.append({"id": f"g-{i + 1}", "maxDrawdown": num(rd.cell(26 + i, 9).value),
                          "troughDate": iso(rd.cell(26 + i, 10).value)})

    # Calendar-year block reads rows up to 2669 in the workbook, so 2026 is partial there: keep 2016-2025.
    calendar = []
    for i in range(5):
        for prefix, col in (("ug", 2), ("g", 8)):
            years = {str(2016 + y): num(rd.cell(34 + y, col + i).value) for y in range(10)}
            calendar.append({"id": f"{prefix}-{i + 1}", "years": years})

    def ef_block(ws, first_asset_row, n_assets, point_rows, weight_col0):
        assets = [ws.cell(first_asset_row + i, 1).value for i in range(n_assets)]
        mu = [num(ws.cell(first_asset_row + i, 2).value) for i in range(n_assets)]
        sigma = [num(ws.cell(first_asset_row + i, 3).value) for i in range(n_assets)]
        cov = [[num(ws.cell(first_asset_row + i, 6 + j).value) for j in range(n_assets)] for i in range(n_assets)]
        points = []
        for r in point_rows:
            points.append({
                "label": ws.cell(r, 1).value,
                "weights": [num(ws.cell(r, weight_col0 + j).value) or 0.0 for j in range(n_assets)],
                "ret": num(ws.cell(r, 2).value), "vol": num(ws.cell(r, 3).value), "sharpe": num(ws.cell(r, 4).value),
            })
        return {"tickers": assets, "mu": mu, "sigma": sigma, "cov": cov, "points": points}

    ef_constrained = ef_block(sheet(wb, "EF (Constrained)"), 6, 7, range(16, 43), 5)
    ef_long_only = ef_block(sheet(wb, "Efficient Frontier"), 6, 13, range(23, 40), 5)

    scen_tickers = [scen.cell(6 + i, 1).value for i in range(17)]
    scenario = {
        "tickers": scen_tickers,
        "classes": [scen.cell(6 + i, 2).value for i in range(17)],
        "weights": [num(scen.cell(6 + i, 4).value) or 0.0 for i in range(17)],
        "mu": [num(summary.cell(4 + i, 7).value) for i in range(17)],
        "sigma": [num(summary.cell(4 + i, 8).value) for i in range(17)],
        "corr": [[num(rd.cell(6 + i, 2 + j).value) for j in range(17)] for i in range(17)],
        "expected": {"ret": num(scen["B34"].value), "vol": num(scen["B35"].value), "sharpe": num(scen["B36"].value)},
        "checks": {
            "total": num(scen["B27"].value), "maxSingle": num(scen["B28"].value),
            "maxClass": num(scen["B29"].value), "core": num(scen["B30"].value),
            "synthetic": num(scen["B31"].value),
        },
    }

    rel = sheet(wb, "Reliability")
    reliability = []
    for r in range(7, 25):
        reliability.append({
            "ticker": rel.cell(r, 1).value,
            "firstActual": iso(rel.cell(r, 2).value),
            "actualDays": num(rel.cell(r, 3).value),
            "backcastDays": num(rel.cell(r, 4).value),
            "actualShare": num(rel.cell(r, 5).value),
            "rating": rel.cell(r, 6).value,
        })

    fixture = {
        "generatedFrom": f"{Path(path).name} recalculated with LibreOffice (headless)",
        "reliability": reliability,
        "tickers": tickers,
        "windows": windows,
        "portfolios": portfolios,
        "correlation": corr,
        "drawdowns": drawdowns,
        "calendarYears": calendar,
        "efficientFrontier": {"constrained": ef_constrained, "longOnly": ef_long_only},
        "scenario": scenario,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(fixture, indent=1) + "\n")
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
